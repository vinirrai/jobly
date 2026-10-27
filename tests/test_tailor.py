from __future__ import annotations

import io
import json
from datetime import date

import anthropic
import httpx2
import pytest
from docx import Document

from jobly.tailor import prompts, render
from jobly.tailor.assemble import assemble_resume
from jobly.tailor.llm import ClaudeLLM, LLMError
from jobly.tailor.pipeline import Tailor
from jobly.tailor.resume_import import import_resume, resume_content
from jobly.tailor.schemas import CoverLetter, JobAnalysis

from .conftest import make_analysis, make_letter, make_plan, make_tailored


def _items(doc):
    return [i for g in doc.skills for i in g.items]


def test_assemble_list_mode_enforces_honesty(profile):
    doc = assemble_resume(profile, make_tailored(), make_analysis())
    # Only real profile entries, deduped, facts copied from the profile
    assert [(x.company, x.title, x.start) for x in doc.experience] == [
        ("Brightline Analytics", "Software Engineering Intern", "Jun 2025")]
    assert doc.experience[0].bullets == ["Built Python ETL into PostgreSQL, cutting reporting time 95%"]
    # Empty tailored bullets fall back to the profile's own bullets
    assert doc.projects[0].bullets == profile.projects[0].bullets
    items = _items(doc)
    assert "Kubernetes" not in items  # stretch skills are never claimed
    assert "Tableau" in items and "A/B testing" in items  # learnable skills listed (mode=list)
    assert "AWS Certified Cloud Practitioner" not in items  # credentials are never auto-added
    assert doc.added_to_learn == ["Tableau", "A/B testing"]
    assert doc.education[0].coursework == ["Databases", "Statistics"]  # invented coursework dropped
    assert doc.name == profile.name and doc.email == profile.email


def test_assemble_learning_and_off_modes(profile):
    profile.preferences.learnable_skills_mode = "learning"
    doc = assemble_resume(profile, make_tailored(), make_analysis())
    assert "Tableau" not in _items(doc)
    assert doc.currently_learning == ["Tableau", "A/B testing"] and doc.added_to_learn == []
    profile.preferences.learnable_skills_mode = "off"
    doc = assemble_resume(profile, make_tailored(), make_analysis())
    assert "Tableau" not in _items(doc) and not doc.currently_learning


def test_assemble_falls_back_when_no_valid_experience(profile):
    t = make_tailored()
    t.experience = []
    doc = assemble_resume(profile, t, make_analysis())
    assert len(doc.experience) == len(profile.experience)


def test_prompts_carry_profile_indices_and_rules(profile, job):
    system = prompts.system_blocks(profile)
    assert "experience[0]" in system[1]["text"] and "projects[0]" in system[1]["text"]
    assert system[1]["cache_control"] == {"type": "ephemeral"}
    assert "preferences" not in system[1]["text"]
    a = make_analysis()
    r = prompts.resume_prompt(profile, job, a)
    assert "Tableau" in r and "Never list these stretch requirements as skills: Kubernetes" in r
    c = prompts.cover_letter_prompt(profile, job, a)
    assert "passion" in c and "grit" in c and "Kubernetes" in c
    lp = prompts.learning_plan_prompt(profile, job, a)
    assert "Tableau" in lp and "Kubernetes" in lp


def test_pipeline_builds_kit(profile, job, fake_llm):
    parts, stages = [], []
    kit = Tailor(fake_llm).build(profile, job, on_part=lambda n, d: parts.append(n), on_stage=stages.append)
    assert parts[0] == "analysis" and set(parts) == {"analysis", "resume", "cover_letter", "learning_plan"}
    assert [c[0] for c in fake_llm.calls][0] == "JobAnalysis" and len(fake_llm.calls) == 4
    assert kit.resume.experience[0].company == "Brightline Analytics"
    assert len(stages) == 2


def test_renderers(profile):
    doc = assemble_resume(profile, make_tailored(), make_analysis())
    md = render.resume_markdown(doc)
    assert md.startswith("# Alex Rivera") and "## Experience" in md and "Tableau" in md
    html = render.resume_html(doc)
    assert "<h2>Skills</h2>" in html and "Brightline Analytics" in html and "<script" not in html
    d = Document(io.BytesIO(render.resume_docx(doc)))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Alex Rivera" in text and "EXPERIENCE" in text and "Brightline" in text
    doc.email = doc.phone = doc.location = ""
    doc.links = {}
    render.resume_docx(doc)  # no contact info must not crash

    letter = make_letter()
    txt = render.cover_letter_text(letter, profile, today=date(2026, 10, 27))
    assert "October 27, 2026" in txt and txt.strip().endswith("Alex Rivera")
    Document(io.BytesIO(render.cover_letter_docx(letter, profile)))
    assert "Dear Hiring Team" in render.cover_letter_html(letter, profile)

    plan_md = render.learning_plan_markdown(make_plan(), "Data Analyst Intern", "Acme")
    assert "Interview sprint" in plan_md and "- [ ] Tableau basics" in plan_md
    assert "search: “tableau free training videos”" in plan_md and "(https://kubernetes.io/docs/home/)" in plan_md


def test_resume_import_content_blocks():
    pdf = resume_content("cv.pdf", b"%PDF-1.4 fake")
    assert pdf[0]["type"] == "document" and pdf[0]["source"]["media_type"] == "application/pdf"
    assert "<resume>" in resume_content("cv.txt", b"Jane Doe")[0]["text"]


def test_import_resume_maps_links():
    class LLM:
        model = "x"

        def structured(self, *, system, messages, schema):
            return schema.model_validate({
                "name": "Jane", "email": "j@x.com", "phone": "", "location": "", "headline": "", "summary": "",
                "links": [{"label": "GitHub", "url": "github.com/j"}, {"label": "blog", "url": ""}],
                "skills": [{"name": "SQL", "level": "advanced", "years": None}], "experience": [], "projects": [],
                "education": [], "certifications": [], "awards": [], "stories": []})

    p = import_resume(LLM(), "cv.txt", b"Jane")
    assert p.links == {"github": "github.com/j"} and p.skills[0].name == "SQL"


# ------------------------------------------------- Claude API wire format

def _sse(message_json: dict) -> bytes:
    text = json.dumps(message_json)
    events = [
        ("message_start", {"type": "message_start", "message": {
            "id": "msg_1", "type": "message", "role": "assistant", "model": "claude-opus-5-5", "content": [],
            "stop_reason": None, "stop_sequence": None,
            "usage": {"input_tokens": 10, "output_tokens": 1, "cache_read_input_tokens": 0}}}),
        ("content_block_start", {"type": "content_block_start", "index": 0, "content_block": {"type": "text", "text": ""}}),
        ("content_block_delta", {"type": "content_block_delta", "index": 0, "delta": {"type": "text_delta", "text": text}}),
        ("content_block_stop", {"type": "content_block_stop", "index": 0}),
        ("message_delta", {"type": "message_delta", "delta": {"stop_reason": "end_turn", "stop_sequence": None},
                           "usage": {"output_tokens": 50}}),
        ("message_stop", {"type": "message_stop"}),
    ]
    return "".join(f"event: {e}\ndata: {json.dumps(d)}\n\n" for e, d in events).encode()


def _client(handler):
    return anthropic.Anthropic(api_key="test-key", max_retries=0,
                               http_client=anthropic.DefaultHttpxClient(transport=httpx2.MockTransport(handler)))


def test_claude_llm_request_shape_and_parsing():
    captured = {}
    letter = make_letter()

    def handler(request):
        captured["body"] = json.loads(request.content)
        captured["headers"] = request.headers
        return httpx2.Response(200, content=_sse(letter.model_dump()), headers={"content-type": "text/event-stream"})

    llm = ClaudeLLM(model="claude-opus-5-5", effort="high", client=_client(handler))
    out = llm.structured(system=[{"type": "text", "text": "sys"}],
                         messages=[{"role": "user", "content": "hi"}], schema=CoverLetter)
    assert out == letter
    body = captured["body"]
    assert body["model"] == "claude-opus-5-5" and body["stream"] is True
    assert body["output_config"]["effort"] == "high"
    assert body["output_config"]["format"]["type"] == "json_schema"
    assert body["fallbacks"] == "default"
    assert "server-side-fallback-2026-07-01" in captured["headers"]["anthropic-beta"]
    assert "thinking" not in body and "temperature" not in body


def test_claude_llm_errors():
    def unauthorized(request):
        return httpx2.Response(401, json={"type": "error", "error": {"type": "authentication_error", "message": "bad"}})

    with pytest.raises(LLMError, match="ANTHROPIC_API_KEY"):
        ClaudeLLM(client=_client(unauthorized)).structured(system=[], messages=[{"role": "user", "content": "x"}],
                                                           schema=JobAnalysis)

    def refusal(request):
        data = _sse({}).decode().replace('"stop_reason": "end_turn"', '"stop_reason": "refusal"')
        return httpx2.Response(200, content=data.encode(), headers={"content-type": "text/event-stream"})

    with pytest.raises(LLMError, match="declined"):
        ClaudeLLM(client=_client(refusal)).structured(system=[], messages=[{"role": "user", "content": "x"}],
                                                      schema=CoverLetter)


def test_no_fallbacks_for_other_models():
    captured = {}

    def handler(request):
        captured["body"] = json.loads(request.content)
        return httpx2.Response(200, content=_sse(make_letter().model_dump()), headers={"content-type": "text/event-stream"})

    ClaudeLLM(model="claude-haiku-5-5", client=_client(handler)).structured(
        system=[], messages=[{"role": "user", "content": "x"}], schema=CoverLetter)
    assert "fallbacks" not in captured["body"]
