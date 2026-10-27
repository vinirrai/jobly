"""Render resumes, cover letters and learning plans as Markdown, print-ready HTML and Word."""

from __future__ import annotations

import html
import io
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from ..models import Profile
from .assemble import ResumeDoc
from .schemas import CoverLetter, LearningPlan

e = html.escape


def _dates(start: str, end: str) -> str:
    if start and end:
        return f"{start} – {end}"
    return start or end


def _contact(doc: ResumeDoc) -> list[str]:
    parts = [p for p in (doc.location, doc.email, doc.phone) if p]
    parts += [v for v in doc.links.values() if v]
    return parts


def _edu_line(ed) -> str:
    deg = ", ".join(p for p in (ed.degree, ed.field) if p)
    return f"{deg}{' — GPA ' + ed.gpa if ed.gpa else ''}"


# ------------------------------------------------------------------ resume

def resume_markdown(doc: ResumeDoc) -> str:
    out = [f"# {doc.name}", " | ".join(_contact(doc)), ""]
    if doc.headline:
        out += [f"**{doc.headline}**", ""]
    if doc.summary:
        out += ["## Summary", doc.summary, ""]
    if doc.skills or doc.currently_learning:
        out.append("## Skills")
        out += [f"- **{g.category}:** {', '.join(g.items)}" for g in doc.skills]
        if doc.currently_learning:
            out.append(f"- **Currently learning:** {', '.join(doc.currently_learning)}")
        out.append("")
    if doc.experience:
        out.append("## Experience")
        for x in doc.experience:
            out.append(f"### {x.title} — {x.company}")
            meta = " | ".join(p for p in (x.location, _dates(x.start, x.end)) if p)
            if meta:
                out.append(f"*{meta}*")
            out += [f"- {b}" for b in x.bullets] + [""]
    if doc.projects:
        out.append("## Projects")
        for p in doc.projects:
            tech = f" ({', '.join(p.tech)})" if p.tech else ""
            link = f" — {p.link}" if p.link else ""
            out.append(f"### {p.name}{tech}{link}")
            out += [f"- {b}" for b in p.bullets] + [""]
    if doc.education:
        out.append("## Education")
        for ed in doc.education:
            out.append(f"### {ed.school}")
            meta = " | ".join(p for p in (_edu_line(ed), ed.location, _dates(ed.start, ed.end)) if p)
            if meta:
                out.append(f"*{meta}*")
            if ed.coursework:
                out.append(f"- Relevant coursework: {', '.join(ed.coursework)}")
            out += [f"- {h}" for h in ed.honors]
            out.append("")
    if doc.certifications:
        out += ["## Certifications"] + [f"- {c}" for c in doc.certifications] + [""]
    if doc.awards:
        out += ["## Awards"] + [f"- {a}" for a in doc.awards] + [""]
    return "\n".join(out).strip() + "\n"


_PRINT_CSS = """
@page { size: letter; margin: 0.5in 0.6in; }
* { box-sizing: border-box; }
body { font-family: Calibri, "Segoe UI", Helvetica, Arial, sans-serif; color: #111; font-size: 10.5pt;
       line-height: 1.32; max-width: 8.5in; margin: 0 auto; padding: 0.4in 0.5in; background: #fff; }
h1 { font-size: 20pt; margin: 0 0 2px; letter-spacing: .3px; }
.contact { color: #333; font-size: 9.5pt; margin-bottom: 4px; }
.headline { font-weight: 600; margin: 4px 0 2px; }
h2 { font-size: 10.5pt; text-transform: uppercase; letter-spacing: 1px; border-bottom: 1px solid #999;
     margin: 10px 0 4px; padding-bottom: 1px; }
.entry { margin-bottom: 6px; break-inside: avoid; }
.row { display: flex; justify-content: space-between; gap: 12px; }
.row b { font-weight: 600; }
.muted { color: #444; font-size: 9.5pt; white-space: nowrap; }
ul { margin: 2px 0 0 0; padding-left: 16px; }
li { margin: 1px 0; }
p { margin: 2px 0; }
.letter p { margin: 0 0 10px; font-size: 11pt; line-height: 1.45; }
@media print { body { padding: 0; } }
"""


def _page(title: str, body: str) -> str:
    return (f"<!doctype html><html lang='en'><head><meta charset='utf-8'><title>{e(title)}</title>"
            f"<meta name='viewport' content='width=device-width, initial-scale=1'>"
            f"<style>{_PRINT_CSS}</style></head><body>{body}</body></html>")


def resume_html(doc: ResumeDoc) -> str:
    h = [f"<h1>{e(doc.name)}</h1>", f"<div class='contact'>{' · '.join(e(c) for c in _contact(doc))}</div>"]
    if doc.headline:
        h.append(f"<div class='headline'>{e(doc.headline)}</div>")
    if doc.summary:
        h.append(f"<h2>Summary</h2><p>{e(doc.summary)}</p>")
    if doc.skills or doc.currently_learning:
        rows = [f"<p><b>{e(g.category)}:</b> {e(', '.join(g.items))}</p>" for g in doc.skills]
        if doc.currently_learning:
            rows.append(f"<p><b>Currently learning:</b> {e(', '.join(doc.currently_learning))}</p>")
        h.append("<h2>Skills</h2>" + "".join(rows))
    if doc.experience:
        h.append("<h2>Experience</h2>")
        for x in doc.experience:
            h.append(
                f"<div class='entry'><div class='row'><span><b>{e(x.title)}</b>, {e(x.company)}"
                f"{' — ' + e(x.location) if x.location else ''}</span>"
                f"<span class='muted'>{e(_dates(x.start, x.end))}</span></div>"
                f"<ul>{''.join(f'<li>{e(b)}</li>' for b in x.bullets)}</ul></div>"
            )
    if doc.projects:
        h.append("<h2>Projects</h2>")
        for p in doc.projects:
            tech = f" <span class='muted'>({e(', '.join(p.tech))})</span>" if p.tech else ""
            link = f" <span class='muted'>{e(p.link)}</span>" if p.link else ""
            h.append(f"<div class='entry'><div><b>{e(p.name)}</b>{tech}{link}</div>"
                     f"<ul>{''.join(f'<li>{e(b)}</li>' for b in p.bullets)}</ul></div>")
    if doc.education:
        h.append("<h2>Education</h2>")
        for ed in doc.education:
            extra = []
            if ed.coursework:
                extra.append(f"<li>Relevant coursework: {e(', '.join(ed.coursework))}</li>")
            extra += [f"<li>{e(x)}</li>" for x in ed.honors]
            h.append(
                f"<div class='entry'><div class='row'><span><b>{e(ed.school)}</b>"
                f"{' — ' + e(_edu_line(ed)) if _edu_line(ed) else ''}</span>"
                f"<span class='muted'>{e(_dates(ed.start, ed.end))}</span></div>"
                f"{'<ul>' + ''.join(extra) + '</ul>' if extra else ''}</div>"
            )
    for title, items in (("Certifications", doc.certifications), ("Awards", doc.awards)):
        if items:
            h.append(f"<h2>{title}</h2><ul>{''.join(f'<li>{e(i)}</li>' for i in items)}</ul>")
    return _page(f"{doc.name} — Resume", "".join(h))


def _set_margins(document: Document, inches: float = 0.6) -> None:
    for s in document.sections:
        s.top_margin = s.bottom_margin = Inches(0.5)
        s.left_margin = s.right_margin = Inches(inches)


def _base_doc(size: float = 10.5) -> Document:
    d = Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(1)
    _set_margins(d)
    return d


def _heading(d: Document, text: str) -> None:
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(10.5)
    # bottom border line
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "999999")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    ppr.append(bdr)


def _bullet(d: Document, text: str) -> None:
    p = d.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(0)
    p.add_run(text)


def _entry_line(d: Document, left_bold: str, left_rest: str, right: str) -> None:
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(7.3), alignment=2)  # right-aligned tab
    r = p.add_run(left_bold)
    r.bold = True
    if left_rest:
        p.add_run(left_rest)
    if right:
        p.add_run("\t" + right).font.color.rgb = RGBColor(0x44, 0x44, 0x44)


def resume_docx(doc: ResumeDoc) -> bytes:
    d = _base_doc()
    name = d.add_paragraph()
    r = name.add_run(doc.name)
    r.bold = True
    r.font.size = Pt(18)
    contact = _contact(doc)
    if contact:
        d.add_paragraph().add_run(" | ".join(contact)).font.size = Pt(9.5)
    if doc.headline:
        d.add_paragraph().add_run(doc.headline).bold = True
    if doc.summary:
        _heading(d, "Summary")
        d.add_paragraph(doc.summary)
    if doc.skills or doc.currently_learning:
        _heading(d, "Skills")
        rows = [(g.category, ", ".join(g.items)) for g in doc.skills]
        if doc.currently_learning:
            rows.append(("Currently learning", ", ".join(doc.currently_learning)))
        for cat, items in rows:
            p = d.add_paragraph()
            p.add_run(f"{cat}: ").bold = True
            p.add_run(items)
    if doc.experience:
        _heading(d, "Experience")
        for x in doc.experience:
            _entry_line(d, x.title, f", {x.company}" + (f" — {x.location}" if x.location else ""),
                        _dates(x.start, x.end))
            for b in x.bullets:
                _bullet(d, b)
    if doc.projects:
        _heading(d, "Projects")
        for p in doc.projects:
            rest = (f" ({', '.join(p.tech)})" if p.tech else "") + (f" — {p.link}" if p.link else "")
            _entry_line(d, p.name, rest, "")
            for b in p.bullets:
                _bullet(d, b)
    if doc.education:
        _heading(d, "Education")
        for ed in doc.education:
            _entry_line(d, ed.school, f" — {_edu_line(ed)}" if _edu_line(ed) else "", _dates(ed.start, ed.end))
            if ed.coursework:
                _bullet(d, "Relevant coursework: " + ", ".join(ed.coursework))
            for h in ed.honors:
                _bullet(d, h)
    for title, items in (("Certifications", doc.certifications), ("Awards", doc.awards)):
        if items:
            _heading(d, title)
            for i in items:
                _bullet(d, i)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------ cover letter

def _long_date(d: date) -> str:
    return f"{d:%B} {d.day}, {d.year}"


def _letter_header(profile: Profile) -> list[str]:
    lines = [profile.name]
    contact = " | ".join(p for p in (profile.email, profile.phone, profile.location) if p)
    if contact:
        lines.append(contact)
    return lines


def cover_letter_text(letter: CoverLetter, profile: Profile, today: date | None = None) -> str:
    today = today or date.today()
    parts = _letter_header(profile) + ["", _long_date(today), "", letter.greeting, ""]
    for para in letter.paragraphs:
        parts += [para, ""]
    parts += [letter.closing, profile.name]
    return "\n".join(parts).strip() + "\n"


def cover_letter_html(letter: CoverLetter, profile: Profile, today: date | None = None) -> str:
    today = today or date.today()
    head = "<br>".join(e(x) for x in _letter_header(profile))
    body = "".join(f"<p>{e(p)}</p>" for p in letter.paragraphs)
    return _page(
        f"{profile.name} — Cover letter",
        f"<div class='letter'><p>{head}</p><p>{_long_date(today)}</p>"
        f"<p>{e(letter.greeting)}</p>{body}<p>{e(letter.closing)}<br>{e(profile.name)}</p></div>",
    )


def cover_letter_docx(letter: CoverLetter, profile: Profile, today: date | None = None) -> bytes:
    today = today or date.today()
    d = _base_doc(11)
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(0.9)
    first = d.add_paragraph()
    first.add_run(profile.name).bold = True
    for line in _letter_header(profile)[1:]:
        d.add_paragraph(line)
    for text in ("", _long_date(today), "", letter.greeting):
        d.add_paragraph(text)
    for para in letter.paragraphs:
        p = d.add_paragraph(para)
        p.paragraph_format.space_after = Pt(10)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    d.add_paragraph(letter.closing)
    d.add_paragraph(profile.name)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------- learning plan

def _resource_md(r) -> str:
    link = r.link.strip()
    if link.lower().startswith("search:"):
        q = link.split(":", 1)[1].strip()
        return f"{r.title} ({r.kind}) — search: “{q}”"
    return f"[{r.title}]({link}) ({r.kind})"


def learning_plan_markdown(plan: LearningPlan, job_title: str = "", company: str = "") -> str:
    target = f" — {job_title}{' at ' + company if company else ''}" if job_title else ""
    out = [f"# Learning plan{target}", "", plan.overview, ""]
    if plan.interview_sprint:
        out += ["## Interview sprint (easy-to-moderate skills)", ""]
        for it in plan.interview_sprint:
            out += [f"### {it.skill} · ~{it.hours}h", f"*Why it matters:* {it.why_it_matters}", "", "**Steps**"]
            out += [f"{i}. {s}" for i, s in enumerate(it.steps, 1)]
            out += ["", "**Resources**"] + [f"- {_resource_md(r)}" for r in it.resources]
            out += ["", f"**Mini project:** {it.mini_project}", "", "**Interview talking points**"]
            out += [f"- {t}" for t in it.talking_points]
            out += ["", f"**You're ready when:** {it.ready_when}", ""]
    if plan.sprint_schedule:
        out += ["## Sprint schedule", ""]
        for day in plan.sprint_schedule:
            out.append(f"**{day.label}**")
            out += [f"- [ ] {t}" for t in day.tasks]
            out.append("")
    if plan.long_term:
        out += ["## Longer-term plan (hard skills)", ""]
        for it in plan.long_term:
            out += [f"### {it.skill} · ~{it.weeks} weeks", f"*How your skills transfer:* {it.transfer_bridge}", "",
                    "**Milestones**"]
            out += [f"- [ ] {m}" for m in it.milestones]
            out += ["", "**Resources**"] + [f"- {_resource_md(r)}" for r in it.resources]
            out += ["", f"**Portfolio project:** {it.portfolio_project}", "",
                    f"**How to talk about it now:** {it.how_to_frame_it_now}", ""]
    if plan.interview_framing:
        out += ["## Showing adaptability and grit in interviews", ""]
        out += [f"- {t}" for t in plan.interview_framing]
    return "\n".join(out).strip() + "\n"
