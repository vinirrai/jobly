"""Create profile.yaml from an existing resume (PDF, Word, or text)."""

from __future__ import annotations

import base64
import io
from pathlib import Path

from docx import Document
from pydantic import BaseModel, Field

from ..models import Education, Experience, Profile, Project, Skill
from . import prompts
from .llm import LLM


class Link(BaseModel):
    label: str = Field(description="e.g. linkedin, github, portfolio")
    url: str


class ProfileDraft(BaseModel):
    name: str
    email: str
    phone: str
    location: str
    links: list[Link]
    headline: str
    summary: str
    skills: list[Skill]
    experience: list[Experience]
    projects: list[Project]
    education: list[Education]
    certifications: list[str]
    awards: list[str]
    stories: list[str]


def _docx_text(data: bytes) -> str:
    d = Document(io.BytesIO(data))
    lines = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text for c in row.cells))
    return "\n".join(lines)


def resume_content(filename: str, data: bytes) -> list[dict]:
    """Message content blocks for the resume file."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return [{
            "type": "document",
            "source": {"type": "base64", "media_type": "application/pdf",
                       "data": base64.standard_b64encode(data).decode()},
        }]
    if suffix == ".docx":
        text = _docx_text(data)
    else:
        text = data.decode("utf-8", errors="replace")
    return [{"type": "text", "text": f"<resume>\n{text}\n</resume>"}]


def import_resume(llm: LLM, filename: str, data: bytes) -> Profile:
    content = resume_content(filename, data) + [{"type": "text", "text": prompts.IMPORT_RESUME}]
    draft = llm.structured(
        system=[{"type": "text", "text": prompts.SYSTEM}],
        messages=[{"role": "user", "content": content}],
        schema=ProfileDraft,
    )
    data_out = draft.model_dump()
    data_out["links"] = {link["label"].lower(): link["url"] for link in data_out["links"] if link["url"]}
    return Profile.model_validate(data_out)
